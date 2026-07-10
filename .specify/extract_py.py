# -*- coding: utf-8 -*-
import os, sys, glob

base = os.getcwd()
outdir = os.path.join(os.environ.get("TEMP",""), "py_offext_out")
os.makedirs(outdir, exist_ok=True)

def out_path(name):
    return os.path.join(outdir, name)

def extract_pptx(path, oname):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation(path)
    sb = ["FILE: %s\n" % os.path.basename(path)]
    for i, slide in enumerate(prs.slides, 1):
        sb.append("\n=== Slide %d ===" % i)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    if not txt and p.text:
                        txt = p.text
                    if txt.strip():
                        sb.append(txt)
            if shape.has_table:
                tbl = shape.table
                sb.append("[TABLE]")
                for row in tbl.rows:
                    cells = [c.text.strip() for c in row.cells]
                    sb.append(" | ".join(cells))
            # group shapes
            if shape.shape_type == 6:  # GROUP
                try:
                    for sub in shape.shapes:
                        if sub.has_text_frame:
                            for p in sub.text_frame.paragraphs:
                                if p.text.strip():
                                    sb.append(p.text)
                except Exception as e:
                    sb.append("[grp err %s]" % e)
        # notes
        try:
            if slide.has_notes_slide:
                nt = slide.notes_slide.notes_text_frame.text
                if nt.strip():
                    sb.append("[NOTES] " + nt)
        except Exception:
            pass
    res = "\n".join(sb)
    with open(out_path(oname), "w", encoding="utf-8") as f:
        f.write(res)
    print("%s <= %s : %d chars" % (oname, os.path.basename(path), len(res)))

def extract_docx(path, oname):
    from docx import Document
    doc = Document(path)
    sb = ["FILE: %s\n" % os.path.basename(path)]
    # iterate body elements in order
    from docx.oxml.ns import qn
    body = doc.element.body
    def iter_block_items(parent):
        for child in parent.iterchildren():
            if child.tag == qn('w:p'):
                yield ('p', child)
            elif child.tag == qn('w:tbl'):
                yield ('tbl', child)
    for kind, el in iter_block_items(body):
        if kind == 'p':
            # find paragraph object
            from docx.text.paragraph import Paragraph
            para = Paragraph(el, doc)
            if para.text.strip():
                sb.append(para.text)
        else:
            from docx.table import Table
            tbl = Table(el, doc)
            sb.append("[TABLE]")
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                sb.append(" | ".join(cells))
    res = "\n".join(sb)
    with open(out_path(oname), "w", encoding="utf-8") as f:
        f.write(res)
    print("%s <= %s : %d chars" % (oname, os.path.basename(path), len(res)))

i = 0
for ext in [".pptx", ".docx"]:
    for f in sorted(glob.glob(os.path.join(base, "*"+ext))):
        i += 1
        oname = "py_%02d.txt" % i
        try:
            if ext == ".pptx":
                extract_pptx(f, oname)
            else:
                extract_docx(f, oname)
        except Exception as e:
            print("%s ERROR: %s" % (os.path.basename(f), e))
print("OUTDIR=%s" % outdir)
